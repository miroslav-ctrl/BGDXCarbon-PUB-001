"""Publish CommonMark as editable Word documents without external assets."""

import logging
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit

from docx import Document
from docx.document import Document as WordDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from markdown_it.token import Token

from .metadata import validate_language, validate_title
from .renderer import _image_as_text, _remove_images, _SafeMarkdown
from .service import PublishingError, PublishingService, PublishResult


@dataclass(frozen=True, slots=True)
class DocxRequest:
    """Single DOCX publication options; title is document metadata."""

    source: Path
    output: Path
    overwrite: bool = False
    title: str | None = None
    language: str = "und"


def _inline(paragraph: Paragraph, tokens: list[Token]) -> None:
    """Retain inline formatting and safe external links in source order."""
    bold = italic = 0
    hyperlink = None
    for token in tokens:
        if token.type == "strong_open":
            bold += 1
        elif token.type == "strong_close":
            bold -= 1
        elif token.type == "em_open":
            italic += 1
        elif token.type == "em_close":
            italic -= 1
        elif token.type == "link_open":
            href = str(token.attrGet("href") or "")
            if urlsplit(href).scheme.lower() not in ("http", "https", "mailto"):
                continue
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(
                qn("r:id"),
                paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True),
            )
            paragraph._p.append(hyperlink)
        elif token.type == "link_close":
            hyperlink = None
        else:
            if not token.content and token.type not in ("softbreak", "hardbreak"):
                continue
            run = paragraph.add_run(
                "\n" if token.type in ("softbreak", "hardbreak") else token.content
            )
            run.bold = bool(bold)
            run.italic = bool(italic)
            if token.type == "code_inline":
                run.font.name = "Consolas"
                run.font.size = Pt(10)
            if hyperlink is not None:
                run.font.color.rgb = RGBColor.from_string("0563C1")
                run.underline = True
                hyperlink.append(run._r)


def _numbering(document: WordDocument, ordered: bool, start: int) -> int:
    """Create independent list numbering, including non-one starting values."""
    root = document.part.numbering_part.element
    abstract_ids = [
        int(e.get(qn("w:abstractNumId"))) for e in root.findall(qn("w:abstractNum"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (
        ("start", str(start)),
        ("numFmt", "decimal" if ordered else "bullet"),
        ("lvlText", "%1." if ordered else "•"),
        ("lvlJc", "left"),
    ):
        child = OxmlElement(f"w:{tag}")
        child.set(qn("w:val"), value)
        level.append(child)
    abstract.append(level)
    root.append(abstract)
    number = root.add_num(abstract_id)
    return int(number.numId)


class DocxRenderer:
    """Render paragraphs, headings, lists, quotations, code and inline markup."""

    def render(self, source: str, title: str, language: str = "und") -> bytes:
        """Return a complete OOXML package, with no macros or embedded images."""
        validate_title(title)
        validate_language(language)
        parser = _SafeMarkdown("commonmark", {"html": False})
        parser.inline.ruler.at("image", _image_as_text)
        tokens = parser.parse(source)
        _remove_images(tokens)
        document = Document()
        document.core_properties.title = title
        document.core_properties.language = language
        document.core_properties.author = ""
        document.core_properties.last_modified_by = ""
        section = document.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2.5)
        section.left_margin = section.right_margin = Cm(2.5)
        for name in ("Normal", "Title", *(f"Heading {i}" for i in range(1, 10))):
            style = document.styles[name]
            style.font.name = "Arial"
            style.font.color.rgb = RGBColor(0, 0, 0)
        normal = document.styles["Normal"]
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing = 1.15
        defaults = document.styles.element.find(qn("w:docDefaults"))
        assert defaults is not None
        properties = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
        lang = properties.find(qn("w:lang"))
        assert lang is not None
        lang.set(qn("w:val"), language)
        lists: list[tuple[int, bool]] = []
        quotes = 0
        paragraph = None
        heading = 0
        for token in tokens:
            if token.type == "heading_open":
                heading = int(token.tag[1:])
            elif token.type == "heading_close":
                heading = 0
            elif token.type == "blockquote_open":
                quotes += 1
            elif token.type == "blockquote_close":
                quotes -= 1
            elif token.type in ("bullet_list_open", "ordered_list_open"):
                lists.append(
                    (
                        _numbering(
                            document,
                            token.type == "ordered_list_open",
                            int(token.attrGet("start") or 1),
                        ),
                        False,
                    )
                )
            elif token.type in ("bullet_list_close", "ordered_list_close"):
                lists.pop()
            elif token.type == "list_item_open":
                lists[-1] = (lists[-1][0], True)
            elif token.type == "paragraph_open":
                paragraph = None
            elif token.type == "inline":
                paragraph = document.add_paragraph(
                    style=f"Heading {heading}" if heading else "Normal"
                )
                paragraph.paragraph_format.left_indent = Cm(0.7 * (len(lists) + quotes))
                if quotes:
                    paragraph.paragraph_format.right_indent = Cm(0.7)
                if lists and lists[-1][1]:
                    num = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                    num.get_or_add_ilvl().val = 0
                    num.get_or_add_numId().val = lists[-1][0]
                    paragraph.paragraph_format.first_line_indent = Cm(-0.4)
                    lists[-1] = (lists[-1][0], False)
                _inline(paragraph, token.children or [])
            elif token.type in ("fence", "code_block"):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.7 * (len(lists) + quotes))
                run = paragraph.add_run(token.content.rstrip("\n"))
                run.font.name = "Consolas"
                run.font.size = Pt(10)
            elif token.type == "hr":
                document.add_paragraph("—")
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()


class DocxPublishingService:
    """Read UTF-8 Markdown and atomically publish its DOCX representation."""

    def __init__(self, logger: logging.Logger) -> None:
        self._logger = logger

    def publish(self, request: DocxRequest) -> PublishResult:
        """Preserve source and previous output on expected publication failure."""
        try:
            source = request.source.resolve(strict=True)
            output = request.output.absolute()
            if source == output.resolve() or (
                output.exists() and source.samefile(output)
            ):
                raise PublishingError("Input and output must be different files.")
            if not source.is_file():
                raise PublishingError("Input must be a regular file.")
            if output.suffix.lower() != ".docx":
                raise PublishingError("DOCX output must have a .docx extension.")
            PublishingService.validate_output(output, request.overwrite)
            content = DocxRenderer().render(
                source.read_text(encoding="utf-8-sig"),
                source.stem if request.title is None else request.title,
                request.language,
            )
            PublishingService.write_bytes(output, content, request.overwrite)
        except (OSError, UnicodeError, ValueError) as error:
            raise PublishingError(f"Unable to publish DOCX: {error}") from error
        self._logger.info("Published DOCX: %s", output)
        return PublishResult(output, len(content))
