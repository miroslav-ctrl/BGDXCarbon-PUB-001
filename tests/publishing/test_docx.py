"""Verify editable OOXML, source fidelity and atomic DOCX publication."""

import logging
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from bgdxpublisher.cli.app import main
from bgdxpublisher.publishing.docx import (
    DocxPublishingService,
    DocxRenderer,
    DocxRequest,
)
from bgdxpublisher.publishing.service import PublishingError, PublishingService


def test_editable_document_preserves_content_and_structure() -> None:
    data = DocxRenderer().render(
        "# Izveštaj Ћирилица\n\n## Section\n\n### Three\n\n"
        "Text **bold** and *italic* with `code`.\nline  \nbreak\n\n"
        "- First\n  - Nested\n- Second\n\n3. Third\n4. Fourth\n\n"
        '> Quote\n>\n> More\n\n```python\nprint("Zdravo")\n```\n\n'
        "    indented\n\n---\n",
        "Izveštaj",
        "sr-Latn",
    )
    doc = Document(BytesIO(data))
    assert doc.core_properties.title == "Izveštaj"
    assert doc.core_properties.language == "sr-Latn"
    assert doc.paragraphs[0].text == "Izveštaj Ћирилица"
    assert [p.style.name for p in doc.paragraphs[:3]] == [
        "Heading 1",
        "Heading 2",
        "Heading 3",
    ]
    runs = doc.paragraphs[3].runs
    assert any(r.text == "bold" and r.bold for r in runs)
    assert any(r.text == "italic" and r.italic for r in runs)
    assert any(r.text == "code" and r.font.name == "Consolas" for r in runs)
    assert "\nline\nbreak" in doc.paragraphs[3].text
    numbered = [p for p in doc.paragraphs if p._p.xpath("./w:pPr/w:numPr")]
    assert len(numbered) == 5
    assert (
        numbered[1].paragraph_format.left_indent
        > numbered[0].paragraph_format.left_indent
    )
    assert doc.part.numbering_part.element.xpath('.//w:start[@w:val="3"]')
    assert 'print("Zdravo")' in [p.text for p in doc.paragraphs]
    assert "indented" in [p.text for p in doc.paragraphs]
    assert doc.paragraphs[-1].text == "—"
    assert doc.paragraphs[-4].paragraph_format.left_indent > 0
    with ZipFile(BytesIO(data)) as package:
        assert b"sr-Latn" in package.read("word/styles.xml")


def test_links_and_images_are_safe_and_visible() -> None:
    data = DocxRenderer().render(
        "[**Safe** link](https://example.com?q=1&b=2) "
        "[mail](mailto:a@example.com) [http](http://example.com) "
        "[local](other.md) [jump](#heading) "
        "[unsafe](javascript:alert(1)) <script>alert(1)</script> "
        "![**Formatted** `alt`](data:image/png;base64,no) "
        "![nested ![inside](x)](https://example.com/image.png)",
        "Test",
    )
    doc = Document(BytesIO(data))
    text = "".join(p.text for p in doc.paragraphs)
    assert "Safe link" in text and "Formatted alt" in text and "nested inside" in text
    assert "<script>alert(1)</script>" in text and "unsafe" in text
    assert "local" in text and "jump" in text
    links = doc.paragraphs[0].hyperlinks
    assert [link.address for link in links] == [
        "https://example.com?q=1&b=2",
        "mailto:a@example.com",
        "http://example.com",
    ]
    assert links[0].runs[0].bold
    with ZipFile(BytesIO(data)) as package:
        assert not any(name.startswith("word/media/") for name in package.namelist())
        assert b"javascript:" not in package.read("word/_rels/document.xml.rels")


def test_independent_lists_and_first_block_markers() -> None:
    doc = Document(
        BytesIO(
            DocxRenderer().render(
                "3. Third\n4. Fourth\n\nSeparate lists\n\n1. First\n2. Second\n\n"
                "- # Heading item\n- ```\n  code item\n  ```\n",
                "Lists",
            )
        )
    )
    paragraphs = [p for p in doc.paragraphs if p._p.xpath("./w:pPr/w:numPr")]
    assert [p.text for p in paragraphs] == [
        "Third",
        "Fourth",
        "First",
        "Second",
        "Heading item",
        "code item",
    ]
    ids = [p._p.pPr.numPr.numId.val for p in paragraphs]
    assert ids[0] == ids[1] and ids[2] == ids[3] and ids[0] != ids[2]
    assert paragraphs[4].style.name == "Heading 1"


@pytest.mark.parametrize(
    "text",
    ["", "# Empty\n\n- \n", "- First\n\n  continuation\n\n- Second\n", "> > Nested\n"],
)
def test_empty_and_compound_blocks(text: str) -> None:
    assert Document(BytesIO(DocxRenderer().render(text, "Test"))) is not None


@pytest.mark.parametrize(
    "field,value", [("title", ""), ("language", "en_US"), ("source", "bad\x01text")]
)
def test_renderer_rejects_invalid_metadata_or_xml(field: str, value: str) -> None:
    kwargs = {"source": "ok", "title": "Test", "language": "und"}
    kwargs[field] = value
    with pytest.raises(ValueError):
        DocxRenderer().render(**kwargs)


def test_cli_publish_overwrite_metadata_and_bom(tmp_path: Path) -> None:
    source, output = tmp_path / "source.md", tmp_path / "out.docx"
    source.write_text("# Čćšžđ Ћирилица", encoding="utf-8-sig")
    before = source.read_bytes()
    command = [
        "publish-docx",
        str(source),
        "--output",
        str(output),
        "--title",
        "Izveštaj",
        "--lang",
        "sr-Cyrl",
    ]
    assert main(command) == 0
    doc = Document(output)
    assert doc.core_properties.title == "Izveštaj"
    assert doc.core_properties.language == "sr-Cyrl"
    assert doc.paragraphs[0].text == "Čćšžđ Ћирилица"
    previous = output.read_bytes()
    assert main(command) == 1
    assert output.read_bytes() == previous
    source.write_text("# New", encoding="utf-8")
    assert main([*command, "--overwrite"]) == 0
    assert Document(output).paragraphs[0].text == "New"
    assert before != source.read_bytes()


@pytest.mark.parametrize(
    "arguments",
    [
        [],
        ["file.md"],
        ["file.md", "--output", "x.docx", "--lang", "en_US"],
        ["file.md", "--output", "x.docx", "--title", ""],
        ["file.md", "--output", "x.docx", "--theme", "dark"],
    ],
)
def test_cli_bad_arguments(arguments: list[str]) -> None:
    with pytest.raises(SystemExit) as error:
        main(["publish-docx", *arguments])
    assert error.value.code == 2


def test_failures_preserve_files(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source, output = tmp_path / "source.md", tmp_path / "out.docx"
    source.write_text("# Test", encoding="utf-8")
    original = source.read_bytes()
    service = DocxPublishingService(logging.getLogger("test"))
    result = service.publish(DocxRequest(source, output))
    assert result.bytes_written == len(output.read_bytes())
    assert Document(output).core_properties.title == "source"
    assert source.read_bytes() == original
    with pytest.raises(PublishingError, match="different files"):
        service.publish(DocxRequest(source, source, True))
    alias = tmp_path / "alias.docx"
    alias.hardlink_to(source)
    with pytest.raises(PublishingError, match="different files"):
        service.publish(DocxRequest(source, alias, True))
    with pytest.raises(PublishingError, match="regular file"):
        service.publish(DocxRequest(tmp_path, output, True))
    with pytest.raises(PublishingError, match="extension"):
        service.publish(DocxRequest(source, tmp_path / "x.html"))
    with pytest.raises(PublishingError):
        service.publish(DocxRequest(tmp_path / "missing", output, True))
    previous = output.read_bytes()
    source.write_bytes(b"\xff")
    with pytest.raises(PublishingError):
        service.publish(DocxRequest(source, output, True))
    assert output.read_bytes() == previous
    source.write_text("# Test", encoding="utf-8")

    def fail(*args: object) -> None:
        raise OSError("commit failed")

    monkeypatch.setattr(PublishingService, "write_bytes", fail)
    with pytest.raises(PublishingError, match="commit failed"):
        service.publish(DocxRequest(source, output, True))
    assert output.read_bytes() == previous
    assert (
        main(["publish-docx", str(source), "--output", str(output), "--overwrite"]) == 1
    )


def test_docx_runtime_closes(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from bgdxpublisher.runtime import PublisherApplication, RuntimeState

    application = PublisherApplication()
    monkeypatch.setattr(
        "bgdxpublisher.cli.publish.PublisherApplication", lambda **kw: application
    )
    source = tmp_path / "source.md"
    source.write_text("Text", encoding="utf-8")
    assert (
        main(["publish-docx", str(source), "--output", str(tmp_path / "out.docx")]) == 0
    )
    assert application.state is RuntimeState.STOPPED
